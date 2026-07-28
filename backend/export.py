"""Styled Excel (and plain CSV) export of a stage grid.

Export is deliberately secondary here — the AG Grid stage view is the primary
surface — but a hiring committee still circulates a spreadsheet, so it has to
look like one: a frozen bold header, colour-coded verdicts, evidence attached
to the cell it justifies, and columns wide enough to read without resizing.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

VERDICT_FILLS = {
    "Meets": PatternFill("solid", fgColor="C6EFCE"),
    "Partial": PatternFill("solid", fgColor="FFEB9C"),
    "No": PatternFill("solid", fgColor="FFC7CE"),
}
VERDICT_FONTS = {
    "Meets": Font(color="006100"),
    "Partial": Font(color="9C5700"),
    "No": Font(color="9C0006"),
}
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MIN_WIDTH = 10
MAX_WIDTH = 60

STAGE_TITLES = {
    "1": "Stage 1 - Minimum Requirements",
    "2": "Stage 2 - Preferred Qualifications",
    "3": "Stage 3 - Interview Candidates",
    "rejected": "Rejected",
}


def stage_title(stage: str) -> str:
    return STAGE_TITLES.get(stage, f"Stage {stage}")


def build_columns(quals: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Header definitions shared by the Excel and CSV writers."""
    columns = [
        {"key": "rank", "label": "Rank"},
        {"key": "name", "label": "Candidate"},
        {"key": "ai", "label": "AI Recommendation"},
        {"key": "required", "label": "Required Met"},
        {"key": "preferred", "label": "Preferred Met"},
    ]
    counters = {"required": 0, "preferred": 0}
    for qual in quals:
        counters[qual["kind"]] += 1
        prefix = "R" if qual["kind"] == "required" else "P"
        columns.append(
            {
                "key": f"qual:{qual['id']}",
                "label": f"{prefix}{counters[qual['kind']]}. {qual['text']}",
            }
        )
    columns.append({"key": "summary", "label": "Summary"})
    columns.append({"key": "files", "label": "Source File(s)"})
    return columns


def _cell_value(row: dict[str, Any], key: str) -> str:
    if key == "rank":
        return str(row.get("rank", ""))
    if key == "name":
        return str(row.get("name", ""))
    if key == "ai":
        if row.get("error"):
            return "Error"
        return "Pass" if row.get("ai_pass") else "Fail"
    if key == "required":
        return f"{row.get('required_met', 0)}/{row.get('required_total', 0)}"
    if key == "preferred":
        return f"{row.get('preferred_met', 0)}/{row.get('preferred_total', 0)}"
    if key == "summary":
        return str(row.get("error") or row.get("summary") or "")
    if key == "files":
        return ", ".join(row.get("source_files") or [])
    if key.startswith("qual:"):
        verdict = (row.get("verdicts") or {}).get(key[5:]) or {}
        return str(verdict.get("verdict") or "")
    return ""


def to_excel(
    screening: dict[str, Any],
    stage: str,
    quals: list[dict[str, Any]],
    rows: list[dict[str, Any]],
) -> bytes:
    columns = build_columns(quals)
    wb = Workbook()
    ws = wb.active
    ws.title = stage_title(stage)[:31]

    title = screening.get("job_title") or "Resume Screening"
    ws["A1"] = f"{title} — {stage_title(stage)}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (
        f"Model: {screening.get('provider') or '—'}/{screening.get('model') or '—'}"
        f"    Candidates: {len(rows)}"
    )
    ws["A2"].font = Font(italic=True, color="595959")

    header_row = 4
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=index, value=column["label"])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for r_index, row in enumerate(rows, start=header_row + 1):
        for c_index, column in enumerate(columns, start=1):
            key = column["key"]
            cell = ws.cell(row=r_index, column=c_index, value=_cell_value(row, key))
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=key in ("summary",))

            if key.startswith("qual:"):
                verdict_info = (row.get("verdicts") or {}).get(key[5:]) or {}
                verdict = verdict_info.get("verdict")
                if verdict in VERDICT_FILLS:
                    cell.fill = VERDICT_FILLS[verdict]
                    cell.font = VERDICT_FONTS[verdict]
                cell.alignment = Alignment(horizontal="center", vertical="center")
                evidence = verdict_info.get("evidence")
                if evidence:
                    # Evidence rides along as a cell note so the grid stays
                    # scannable but the justification is never more than a
                    # hover away — the same affordance the AG Grid view gives.
                    cell.comment = Comment(evidence[:2000], "Resume Screening")
            elif key == "ai":
                if row.get("error"):
                    cell.fill = FAIL_FILL
                else:
                    cell.fill = PASS_FILL if row.get("ai_pass") else FAIL_FILL

    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(columns))}{header_row + len(rows)}"
    )
    _size_columns(ws, columns, rows, header_row)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _size_columns(ws, columns: list[dict[str, str]], rows: list[dict[str, Any]], header_row: int) -> None:
    for index, column in enumerate(columns, start=1):
        key = column["key"]
        if key.startswith("qual:"):
            # Verdict columns hold one short word; the long header wraps above.
            width = 14.0
        else:
            longest = max(
                [len(column["label"])] + [len(_cell_value(row, key)) for row in rows]
            )
            width = min(MAX_WIDTH, max(MIN_WIDTH, longest + 2))
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[header_row].height = 46


def to_csv(
    screening: dict[str, Any],
    stage: str,
    quals: list[dict[str, Any]],
    rows: list[dict[str, Any]],
) -> bytes:
    del screening, stage  # the CSV is the grid, without the report chrome
    columns = build_columns(quals)
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow([c["label"] for c in columns])
    for row in rows:
        writer.writerow([_cell_value(row, c["key"]) for c in columns])
    return buffer.getvalue().encode("utf-8-sig")

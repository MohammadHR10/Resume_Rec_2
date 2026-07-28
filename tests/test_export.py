"""Stage export: column layout, verdict colouring, evidence, CSV parity."""

from __future__ import annotations

import io

from openpyxl import load_workbook

from backend import export

SCREENING = {"job_title": "Software Engineer II", "provider": "ulproxy", "model": "gpt-5.2"}

QUALS = [
    {"id": "r1", "text": "Bachelor's degree in Computer Science", "kind": "required"},
    {"id": "r2", "text": "Three years of Python", "kind": "required"},
    {"id": "p1", "text": "Distributed systems", "kind": "preferred"},
]

ROWS = [
    {
        "id": "a",
        "rank": 1,
        "name": "Alice",
        "required_met": 2,
        "required_total": 2,
        "preferred_met": 1,
        "preferred_total": 1,
        "ai_pass": True,
        "summary": "Strong backend engineer.",
        "error": "",
        "source_files": ["alice.pdf"],
        "verdicts": {
            "r1": {"verdict": "Meets", "evidence": "BS Computer Science, 2018"},
            "r2": {"verdict": "Meets", "evidence": "4 years Python"},
            "p1": {"verdict": "Meets", "evidence": "Kafka pipelines"},
        },
    },
    {
        "id": "b",
        "rank": 2,
        "name": "Bob",
        "required_met": 1,
        "required_total": 2,
        "preferred_met": 0,
        "preferred_total": 1,
        "ai_pass": False,
        "summary": "",
        "error": "",
        "source_files": ["bob.pdf"],
        "verdicts": {
            "r1": {"verdict": "Meets", "evidence": "BS CS"},
            "r2": {"verdict": "Partial", "evidence": "2 years Python"},
            "p1": {"verdict": "No", "evidence": ""},
        },
    },
]


def _sheet(stage: str = "1"):
    data = export.to_excel(SCREENING, stage, QUALS, ROWS)
    return load_workbook(io.BytesIO(data)).active


def test_columns_cover_the_rollups_and_one_column_per_qualification():
    labels = [column["label"] for column in export.build_columns(QUALS)]
    assert labels[:5] == ["Rank", "Candidate", "AI Recommendation", "Required Met", "Preferred Met"]
    assert labels[5].startswith("R1. Bachelor's degree")
    assert labels[7].startswith("P1. Distributed systems")
    assert labels[-2:] == ["Summary", "Source File(s)"]


def test_the_header_row_is_bold_filled_and_frozen():
    sheet = _sheet()
    header = sheet.cell(row=4, column=1)
    assert header.value == "Rank"
    assert header.font.bold is True
    assert header.fill.fgColor.rgb.endswith("1F3864")
    assert sheet.freeze_panes == "C5"
    assert sheet.auto_filter.ref.startswith("A4")


def test_the_title_row_names_the_stage_and_the_model():
    sheet = _sheet("2")
    assert sheet["A1"].value == "Software Engineer II — Stage 2 - Preferred Qualifications"
    assert "ulproxy/gpt-5.2" in sheet["A2"].value
    assert sheet.title == "Stage 2 - Preferred Qualification"[:31]


def test_verdict_cells_are_colour_coded_and_carry_their_evidence():
    sheet = _sheet()
    meets = sheet.cell(row=5, column=6)
    assert meets.value == "Meets"
    assert meets.fill.fgColor.rgb.endswith("C6EFCE")
    assert "BS Computer Science" in meets.comment.text

    partial = sheet.cell(row=6, column=7)
    assert partial.value == "Partial"
    assert partial.fill.fgColor.rgb.endswith("FFEB9C")

    absent = sheet.cell(row=6, column=8)
    assert absent.value == "No"
    assert absent.comment is None  # nothing to evidence


def test_the_rollup_and_recommendation_cells_read_as_the_grid_does():
    sheet = _sheet()
    assert sheet.cell(row=5, column=3).value == "Pass"
    assert sheet.cell(row=5, column=4).value == "2/2"
    assert sheet.cell(row=6, column=3).value == "Fail"
    assert sheet.cell(row=6, column=5).value == "0/1"


def test_a_failed_evaluation_is_shown_as_an_error_not_as_a_pass():
    rows = [{**ROWS[0], "error": "provider timed out", "ai_pass": False}]
    sheet = load_workbook(io.BytesIO(export.to_excel(SCREENING, "1", QUALS, rows))).active
    assert sheet.cell(row=5, column=3).value == "Error"
    assert sheet.cell(row=5, column=9).value == "provider timed out"


def test_csv_matches_the_excel_columns():
    text = export.to_csv(SCREENING, "1", QUALS, ROWS).decode("utf-8-sig")
    lines = text.strip().splitlines()
    assert len(lines) == 3
    assert lines[0].split(",")[:3] == ["Rank", "Candidate", "AI Recommendation"]
    assert lines[1].startswith("1,Alice,Pass,2/2,1/1,Meets,Meets,Meets")


def test_an_empty_stage_still_exports_a_valid_sheet():
    sheet = load_workbook(io.BytesIO(export.to_excel(SCREENING, "3", QUALS, []))).active
    assert sheet.cell(row=4, column=1).value == "Rank"
    assert sheet.cell(row=5, column=1).value is None
